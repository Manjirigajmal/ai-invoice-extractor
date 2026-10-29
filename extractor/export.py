"""Write results to a formatted Excel workbook: Invoices, Line Items, Issues."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .schema import InvoiceResult

FILL = {"OK": "E6F4EA", "Check": "FFF4CE", "Needs review": "FDE2E1", "Failed": "E0E0E0"}
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
MONEY = '#,##0.00'

INVOICE_COLS = [
    ("Status", 14), ("Source File", 24), ("Supplier", 30), ("Supplier GSTIN", 18), ("Buyer", 28),
    ("Buyer GSTIN", 18), ("Invoice No", 18), ("Invoice Date", 13), ("Lines", 7), ("Subtotal", 13),
    ("CGST", 11), ("SGST", 11), ("IGST", 11), ("Total", 13), ("Issues", 60),
]


def invoice_rows(results: list[InvoiceResult]) -> list[list]:
    rows = []
    for r in results:
        d = r.data
        issues = r.error or "; ".join(i.message for i in r.issues)
        rows.append([r.status, r.source_file, d.vendor_name, d.vendor_gstin, d.buyer_name, d.buyer_gstin,
                     d.invoice_number, d.invoice_date, len(d.line_items), d.subtotal, d.cgst, d.sgst, d.igst,
                     d.total, issues])
    return rows


def _sheet(ws, columns, rows, money_cols=(), status_col=None):
    ws.append([c for c, _ in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    for i, (_, width) in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for idx in money_cols:
            row[idx].number_format = MONEY
        if status_col is not None:
            colour = FILL.get(row[status_col].value)
            if colour:
                row[status_col].fill = PatternFill("solid", fgColor=colour)
        row[-1].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def to_excel(results: list[InvoiceResult]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"
    _sheet(ws, INVOICE_COLS, invoice_rows(results), money_cols=(9, 10, 11, 12, 13), status_col=0)

    items = wb.create_sheet("Line Items")
    item_rows = [[r.source_file, r.data.invoice_number, n, li.description, li.hsn_sac, li.quantity,
                  li.unit_price, li.amount]
                 for r in results for n, li in enumerate(r.data.line_items, 1)]
    _sheet(items, [("Source File", 24), ("Invoice No", 18), ("Line", 6), ("Description", 44), ("HSN/SAC", 10),
                   ("Qty", 8), ("Unit Price", 13), ("Amount", 14)], item_rows, money_cols=(6, 7))

    probs = wb.create_sheet("Issues")
    issue_rows = [[r.source_file, r.data.invoice_number, i.severity.title(), i.code, i.message]
                  for r in results for i in r.issues]
    issue_rows += [[r.source_file, None, "Error", "extraction_failed", r.error] for r in results if r.error]
    _sheet(probs, [("Source File", 24), ("Invoice No", 18), ("Severity", 10), ("Check", 20), ("Details", 80)],
           issue_rows)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
