import io
import json
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

import pytest
from openpyxl import load_workbook

from extractor import gstin as G
from extractor.export import to_excel
from extractor.pipeline import extract_one, load_folder, run_batch
from extractor.schema import InvoiceData, LineItem
from extractor.validate import check_invoice

ROOT = Path(__file__).resolve().parent.parent
SAMPLES = ROOT / "samples" / "invoices"
MH = G.make("27", "AAHCM4821K")
KA = G.make("29", "ABDFB9051R")


def invoice(**kw) -> InvoiceData:
    base = dict(vendor_name="Vendor", vendor_gstin=MH, buyer_name="Buyer", buyer_gstin=G.make("27", "AAKCS7310P"),
                invoice_number="INV-1", invoice_date="2026-09-01",
                line_items=[LineItem(description="A", quantity=2, unit_price=500, amount=1000)],
                subtotal=1000, cgst=90, sgst=90, igst=0, total=1180)
    base.update(kw)
    return InvoiceData(**base)


def codes(d: InvoiceData) -> set[str]:
    return {i.code for i in check_invoice(d)}


# ---------- GSTIN ----------
def test_gstin_checksum():
    assert G.is_valid(MH)
    assert not G.is_valid(MH[:-1] + ("A" if MH[-1] != "A" else "B"))
    assert not G.is_valid("27ABC")
    assert G.state_code(KA) == "29"


# ---------- validation ----------
def test_clean_invoice_has_no_issues():
    assert codes(invoice()) == set()


def test_total_mismatch():
    assert "total_mismatch" in codes(invoice(total=2180))


def test_line_item_and_subtotal_mismatch():
    d = invoice(line_items=[LineItem(description="A", quantity=2, unit_price=500, amount=1500)])
    assert {"line_item_mismatch", "subtotal_mismatch"} <= codes(d)


def test_inter_state_must_use_igst():
    assert "wrong_tax_type" in codes(invoice(vendor_gstin=KA))
    assert "wrong_tax_type" not in codes(invoice(vendor_gstin=KA, cgst=0, sgst=0, igst=180))


def test_intra_state_must_not_use_igst():
    assert "wrong_tax_type" in codes(invoice(cgst=0, sgst=0, igst=180))


def test_missing_fields_and_bad_gstin():
    c = codes(invoice(invoice_number=None, vendor_gstin="27AAHCM4821K1Z0"))
    assert {"missing_field", "invalid_gstin"} <= c


def test_rounding_tolerance():
    assert codes(invoice(total=1180.60)) == set()


# ---------- full sample batch (offline engine) ----------
def test_offline_batch_catches_every_planted_problem():
    truth = {t["file"]: t for t in json.loads((ROOT / "samples" / "ground_truth.json").read_text())}
    results = run_batch(load_folder(SAMPLES), "offline")
    assert len(results) == 12
    for r in results:
        problem = truth[r.source_file]["planted_problem"]
        error_codes = {i.code for i in r.issues if i.severity == "error"}
        if problem:
            assert problem in error_codes, r.source_file
        else:
            assert r.status == "OK", (r.source_file, r.issues)


def test_excel_export_has_three_sheets():
    results = run_batch(load_folder(SAMPLES), "offline")
    wb = load_workbook(io.BytesIO(to_excel(results)))
    assert wb.sheetnames == ["Invoices", "Line Items", "Issues"]
    assert wb["Invoices"].max_row == 13
    assert wb["Line Items"].max_row == 1 + 24


def test_bad_file_does_not_stop_batch():
    r = extract_one("broken.pdf", b"not a pdf", "offline")
    assert r.status == "Failed" and r.error


# ---------- LLM engines (mocked: no API key or network needed) ----------
LLM_REPLY = invoice().model_dump()


def test_claude_engine_parses_tool_output():
    fake = mock.MagicMock()
    fake.messages.create.return_value = SimpleNamespace(content=[SimpleNamespace(type="tool_use", input=LLM_REPLY)])
    with mock.patch("anthropic.Anthropic", return_value=fake):
        r = extract_one("SGP-2026-0142.pdf", (SAMPLES / "SGP-2026-0142.pdf").read_bytes(), "claude")
    assert r.error is None and r.data.total == 1180
    kwargs = fake.messages.create.call_args.kwargs
    assert kwargs["tool_choice"] == {"type": "tool", "name": "record_invoice"}
    assert "Invoice text" in kwargs["messages"][0]["content"][0]["text"]  # text layer was used


def test_claude_engine_sends_pdf_when_scanned():
    from extractor.llm import extract_with_claude
    fake = mock.MagicMock()
    fake.messages.create.return_value = SimpleNamespace(content=[SimpleNamespace(type="tool_use", input=LLM_REPLY)])
    with mock.patch("anthropic.Anthropic", return_value=fake):
        extract_with_claude(None, b"%PDF-1.4 fake")
    block = fake.messages.create.call_args.kwargs["messages"][0]["content"][0]
    assert block["type"] == "document" and block["source"]["media_type"] == "application/pdf"


def test_gemini_engine_parses_json():
    fake = mock.MagicMock()
    fake.models.generate_content.return_value = SimpleNamespace(parsed=None, text=json.dumps(LLM_REPLY))
    with mock.patch("google.genai.Client", return_value=fake):
        r = extract_one("KOS-0605.pdf", (SAMPLES / "KOS-0605.pdf").read_bytes(), "gemini")
    assert r.error is None and r.data.invoice_number == "INV-1"


def test_missing_api_key_is_reported_not_raised(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.delenv("ANTHROPIC_AUTH_TOKEN", raising=False)
    r = extract_one("x.pdf", (SAMPLES / "KOS-0605.pdf").read_bytes(), "claude")
    assert r.status == "Failed"


@pytest.mark.parametrize("raw,expected", [("03-09-2026", "2026-09-03"), ("09 Sep 2026", "2026-09-09"),
                                          ("12/09/2026", "2026-09-12"), ("2026-09-01", "2026-09-01")])
def test_offline_date_formats(raw, expected):
    from extractor.rules import _date
    assert _date(raw) == expected
